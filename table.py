import os
import django
from openpyxl import load_workbook

# Setup Django env
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'voiceai.settings')
django.setup()

from chat.models import Products  # Replace with your actual app name


def import_products_from_excel_nb(xlsx_file_path):
    wb = load_workbook(filename=xlsx_file_path)
    sheet = wb.active  # Assuming data is in the first sheet

    # Get headers from first row
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]

    products_to_create = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))

        product = Products(
            category=row_dict.get('Category'),
            subcategory=row_dict.get('Subcategory'),
            product_name=row_dict.get('Product Name'),
            price=float(str(row_dict.get('Price')).replace('$','').strip()) if row_dict.get('Price') else 0,
            pregnancy=True if str(row_dict.get('Pregnancy')).strip().lower() == 'yes' else False,
            orthodontics=True if str(row_dict.get('Orthodontics')).strip().lower() == 'yes' else False,
            teething_to_24_months=True if str(row_dict.get('Teething to 24 months')).strip().lower() == 'yes' else False,
            age_2_to_5=True if str(row_dict.get('2 to 5')).strip().lower() == 'yes' else False,
            age_6_to_12=True if str(row_dict.get('6 to 12')).strip().lower() == 'yes' else False,
            age_13_and_above=True if str(row_dict.get('13 to above')).strip().lower() == 'yes' else False,
            description=row_dict.get('Description', '')
        )
        products_to_create.append(product)

    Products.objects.bulk_create(products_to_create)
    print(f"Successfully imported {len(products_to_create)} products from Excel.")


if __name__ == '__main__':
    excel_path = 'products.xlsx'  # Replace with your Excel file path
    import_products_from_excel_nb(excel_path)
